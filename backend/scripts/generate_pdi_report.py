"""Create the NBE Motors Pre-Dispatch Inspection (PDI) Report workbook.

The sheet is editable. Observation cells accept only OK or NOT OK from a dropdown.

    python scripts/generate_pdi_report.py
    python scripts/generate_pdi_report.py /path/to/PDI_Report.xlsx
"""

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

NAVY = "1F4E78"
ZEBRA = "F9FAFB"
LABEL_FILL = "E8EEF4"
SECTION_FILL = "F3F6FA"
WHITE = "FFFFFF"
GRID = "D1D5DB"
INPUT_FILL = "FFFFFF"

LAST_COL = 6
HEADER_ROW = 11
FIRST_DATA_ROW = 12

COLUMN_WIDTHS = {
    "A": 18,
    "B": 45,
    "C": 30,
    "D": 15,
    "E": 28,
    "F": 18,
}

THIN = Border(
    left=Side(style="thin", color=GRID),
    right=Side(style="thin", color=GRID),
    top=Side(style="thin", color=GRID),
    bottom=Side(style="thin", color=GRID),
)

SECTIONS = [
    {
        "title": "1. ELECTRICAL TESTING CONFIRMATION",
        "items": [
            ("Insulation Resistance Test", "", "100%", "TEST Bed"),
            ("Measurement of Resistance of Windings of Stator", "", "100%", "TEST Bed"),
            ("No Load Test", "", "100%", "TEST Bed"),
            ("High Voltage Test", "", "100%", "TEST Bed"),
            (
                "Locked Rotor Readings of Voltage, Current and Power Input at a suitable reduced voltage",
                "",
                "100%",
                "TEST Bed",
            ),
            ("Dimensions", "", "100%", "TEST Bed"),
            ("Earthing", "", "100%", "TEST Bed"),
            ("Terminal Marking", "", "100%", "TEST Bed"),
            (
                "Temperature Rise Test",
                "One Motors of each type & design, manufactured in three months",
                "",
                "TEST Bed",
            ),
            (
                "Full Load Test to determine Efficiency, Power Factor and Slip",
                "One Motors of each type & design, manufactured in three months",
                "",
                "TEST Bed",
            ),
        ],
    },
    {
        "title": "2. COOLING FAN COVER POSITION",
        "items": [
            ("MOTOR ROTATION", "", "100%", "VISUAL"),
        ],
    },
    {
        "title": "3. VERIFICATION ASSEMBLY",
        "items": [
            ("FRAME SIZE", "", "100%", "REF. CHART"),
            ("BEARING Make & SIZE", "", "100%", "As per validation"),
            ("VARNISH Make of Varnish", "", "Batch", "As per validation"),
            (
                "Fitment of Fan Cover, bolts, studs, nuts, Eye bolts, lugs, plugs, flanges",
                "",
                "100%",
                "VISUAL",
            ),
            ("EARTH PLATE", "", "BATCH", "VISUAL"),
        ],
    },
    {
        "title": "4. PAINTING VERIFICATION",
        "items": [
            ("PAINT COLOR SHADE Casting-Std", "", "100%", "Approved TEMPLATE"),
            ("CED Coating thickness", "20 to 40 Micron", "", "DFT meter"),
            ("Free from VISUAL DEFECTS (Painting-damage, rust etc)", "", "100%", "VISUAL"),
        ],
    },
    {
        "title": "5. FINAL PRODUCT VERIFICATION",
        "items": [
            ("DIRECTION OF ROTATION", "", "100%", "VISUAL"),
            ("NAME PLATE MODEL DETAILS", "", "100%", "VISUAL"),
            (
                "KIRLOSKAR NAME & ADDRESS VERIFICATION No Spelling mistake & legibility issue",
                "",
                "100%",
                "VISUAL",
            ),
            ("All over body of dents, damages, dirt, rust & scratch Free", "", "100%", "VISUAL"),
            ("BRAND LOGO AVAILABLE ON MOTORS", "", "100%", "VISUAL"),
            (
                "BRAND LOGO APPEARANCE - Damage/broken, excess shot blasting, legibility issue not allowed. Fan cover bolt plating to be ensured",
                "",
                "100%",
                "VISUAL",
            ),
            (
                "MOTOR BODY CASTING - No fin mismatch, No sand drop between fins, Mounting hole burr not allowed.",
                "",
                "100%",
                "VISUAL",
            ),
            (
                "CASTING FINISH - Blow hole, cold shut, crack, Sand Drop & Poor shot blasting not allowed",
                "",
                "100%",
                "VISUAL",
            ),
            ("Rear side Water Removal Pressurised Air", "", "100%", "VISUAL"),
            (
                "Sound Testing for Pump Body and Volute (sound testing-free hang the casting & bang by steel rod)",
                "",
                "100%",
                "VISUAL",
            ),
            ("Terminal Box Cover as per Brand LOGO", "", "100%", "VISUAL"),
        ],
    },
    {
        "title": "6. PACKAGING",
        "items": [
            ("WARRANTY CARD", "", "100%", "VISUAL"),
            ("INSTRUCTION MANUAL", "", "100%", "VISUAL"),
            ("OUTER BOX MRP DETAILS", "", "100%", "VISUAL"),
            ("MASTER BOX PLY", "", "100%", "No. of Ply"),
            ("INNER POLYTHENE COVER PROPERLY SEALED / NO OPEN END", "", "100%", "VISUAL"),
            ("PART NUMBER VERIFICATION", "", "100%", "VISUAL COMPARE WITH PO"),
            ("NET WEIGHT", "", "ACTUAL", "ACTUAL-Value to be mentioned in Kg"),
            ("GROSS WEIGHT", "", "ACTUAL", "ACTUAL-Value to be mentioned in Kg"),
            (
                "Material test report for all Casting parts & Forging parts",
                "",
                "Per Heat/Batch Code",
                "Attach actual test report as per Material Grade",
            ),
        ],
    },
]

METADATA = [
    ("REV NO / DATE", "01/04/2025", "PART NO", ""),
    ("PART DESCRIPTION", "", "HP/KW", ""),
    ("CUSTOMER NAME", "", "PUMP PIPE SIZE", ""),
    ("NAME PLATE SERIAL NUMBER", "", "BATCH NO", ""),
    ("INVOICE NO", "", "PO NO", ""),
    ("INVOICE DATE", "", "PO QTY", ""),
]


def _fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)


def _estimate_lines(text, width):
    if not text:
        return 1
    capacity = max(int(width) - 1, 8)
    lines = 0
    for paragraph in str(text).split("\n"):
        words = paragraph.split()
        if not words:
            lines += 1
            continue
        length = 0
        count = 1
        for word in words:
            extra = len(word) if length == 0 else len(word) + 1
            if length + extra > capacity:
                count += 1
                length = len(word)
            else:
                length += extra
        lines += count
    return lines


def _apply_border(ws, row, start_col=1, end_col=LAST_COL):
    for col in range(start_col, end_col + 1):
        ws.cell(row, col).border = THIN


def _write_banner(ws):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=LAST_COL)
    company = ws.cell(1, 1, "NBE MOTORS PRIVATE LIMITED")
    company.font = Font(name="Calibri", bold=True, size=16, color=NAVY)
    company.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=LAST_COL)
    subtitle = ws.cell(2, 1, "PRE DISPATCH INSPECTION REPORT")
    subtitle.font = Font(name="Calibri", bold=True, size=11, color=NAVY)
    subtitle.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18

    for col in range(1, LAST_COL + 1):
        ws.cell(1, col).border = THIN
        ws.cell(2, col).border = THIN


def _write_metadata(ws):
    label_font = Font(name="Calibri", bold=True, size=10, color=NAVY)
    value_font = Font(name="Calibri", size=10)
    label_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    value_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    for offset, (left_label, left_value, right_label, right_value) in enumerate(METADATA):
        row = 4 + offset
        ws.row_dimensions[row].height = 22

        ws.cell(row, 1, left_label)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
        ws.cell(row, 2, left_value)

        ws.cell(row, 4, right_label)
        ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=6)
        ws.cell(row, 5, right_value)

        for col in (1, 4):
            cell = ws.cell(row, col)
            cell.font = label_font
            cell.fill = _fill(LABEL_FILL)
            cell.alignment = label_alignment
        for col in (2, 5):
            cell = ws.cell(row, col)
            cell.font = value_font
            cell.fill = _fill(INPUT_FILL)
            cell.alignment = value_alignment
            cell.number_format = "@"

        _apply_border(ws, row)


def _write_table_header(ws):
    headers = [
        "SR NO",
        "PARAMETER",
        "SPECIFICATION",
        "FREQ",
        "INSPECTION METHOD",
        "OBSERVATION",
    ]
    header_font = Font(name="Calibri", bold=True, size=10, color=WHITE)
    header_fill = _fill(NAVY)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(HEADER_ROW, col, title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = THIN
    ws.row_dimensions[HEADER_ROW].height = 24


def _write_sections(ws):
    body_font = Font(name="Calibri", size=10)
    section_font = Font(name="Calibri", bold=True, size=9, color=NAVY)
    wrap_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    wrap_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    section_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row = FIRST_DATA_ROW
    data_index = 0
    for section in SECTIONS:
        start = row
        for parameter, specification, freq, method in section["items"]:
            zebra = _fill(ZEBRA) if data_index % 2 else _fill(WHITE)
            data_index += 1
            values = {
                2: parameter,
                3: specification,
                4: freq,
                5: method,
                6: "",
            }
            alignments = {
                2: wrap_left,
                3: wrap_left,
                4: wrap_center,
                5: wrap_left,
                6: wrap_center,
            }
            for col, value in values.items():
                cell = ws.cell(row, col, value)
                cell.font = body_font
                cell.fill = zebra
                cell.alignment = alignments[col]
                cell.border = THIN

            lines = max(
                _estimate_lines(parameter, COLUMN_WIDTHS["B"]),
                _estimate_lines(specification, COLUMN_WIDTHS["C"]),
                _estimate_lines(method, COLUMN_WIDTHS["E"]),
            )
            ws.row_dimensions[row].height = max(20, 15 * lines + 6)
            row += 1

        end = row - 1
        if end > start:
            ws.merge_cells(start_row=start, start_column=1, end_row=end, end_column=1)
        section_cell = ws.cell(start, 1, section["title"])
        section_cell.font = section_font
        section_cell.alignment = section_alignment
        section_fill = _fill(SECTION_FILL)
        for section_row in range(start, end + 1):
            cell = ws.cell(section_row, 1)
            cell.fill = section_fill
            cell.border = THIN
            cell.alignment = section_alignment

    return row - 1


def _write_footer(ws, last_data_row):
    remarks_label_row = last_data_row + 2
    remarks_start = remarks_label_row + 1
    remarks_end = remarks_start + 2
    signoff_row = remarks_end + 2

    label_font = Font(name="Calibri", bold=True, size=10, color=NAVY)
    body_font = Font(name="Calibri", size=10)

    ws.merge_cells(
        start_row=remarks_label_row,
        start_column=1,
        end_row=remarks_label_row,
        end_column=LAST_COL,
    )
    label = ws.cell(remarks_label_row, 1, "Remarks")
    label.font = label_font
    label.fill = _fill(LABEL_FILL)
    label.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[remarks_label_row].height = 18
    _apply_border(ws, remarks_label_row)

    ws.merge_cells(
        start_row=remarks_start,
        start_column=1,
        end_row=remarks_end,
        end_column=LAST_COL,
    )
    remarks = ws.cell(remarks_start, 1, "")
    remarks.font = body_font
    remarks.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    for row in range(remarks_start, remarks_end + 1):
        ws.row_dimensions[row].height = 18
        for col in range(1, LAST_COL + 1):
            cell = ws.cell(row, col)
            cell.border = THIN
            cell.fill = _fill(WHITE)

    ws.merge_cells(start_row=signoff_row, start_column=1, end_row=signoff_row, end_column=3)
    ws.merge_cells(start_row=signoff_row, start_column=4, end_row=signoff_row, end_column=6)
    checked = ws.cell(signoff_row, 1, "Checked by: _____________")
    approved = ws.cell(signoff_row, 4, "Approved by: _____________")
    sign_alignment = Alignment(horizontal="left", vertical="center")
    for cell in (checked, approved):
        cell.font = Font(name="Calibri", bold=True, size=10)
        cell.alignment = sign_alignment
    ws.row_dimensions[signoff_row].height = 28
    _apply_border(ws, signoff_row)
    return signoff_row


def _add_observation_validation(ws, last_data_row):
    validation = DataValidation(
        type="list",
        formula1='"OK,NOT OK"',
        allow_blank=True,
        showErrorMessage=True,
        errorStyle="stop",
        errorTitle="Invalid Entry",
        error="Please select OK or NOT OK from the dropdown list.",
        # openpyxl maps this to Excel's hideDropDown flag. False keeps the dropdown visible.
        showDropDown=False,
    )
    validation.add(f"F{FIRST_DATA_ROW}:F{last_data_row}")
    ws.add_data_validation(validation)


def build_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "PDI Report"

    for letter, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[letter].width = width

    _write_banner(ws)
    _write_metadata(ws)
    _write_table_header(ws)
    last_data_row = _write_sections(ws)
    _write_footer(ws, last_data_row)
    _add_observation_validation(ws, last_data_row)

    ws.freeze_panes = "A12"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.horizontalCentered = True
    ws.print_title_rows = "1:11"
    ws.page_margins.left = 0.4
    ws.page_margins.right = 0.4
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5
    ws.sheet_view.showGridLines = False

    wb.properties.title = "PRE DISPATCH INSPECTION REPORT"
    wb.properties.subject = "NBE MOTORS PRIVATE LIMITED"
    return wb


def default_output_path():
    return Path(__file__).resolve().parents[1] / "templates" / "NBE_Pre_Dispatch_Inspection_Report.xlsx"


def main(argv=None):
    args = list(sys.argv[1:] if argv is None else argv)
    output = Path(args[0]).expanduser() if args else default_output_path()
    output.parent.mkdir(parents=True, exist_ok=True)
    build_workbook().save(output)
    print(f"Wrote {output}")
    return output


if __name__ == "__main__":
    main()
